from datetime import datetime
from pathlib import Path

from django.http import HttpResponse
from openpyxl.drawing.image import Image as OpenPyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="8B7355", end_color="6F5B44", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="F5F1E8", end_color="E8DCC8", fill_type="solid")
WHITE_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=18, bold=True, color="5D4037")
SUBTITLE_FONT = Font(size=11, color="6F5B44")
META_FONT = Font(size=9, color="8B7355")
DATA_FONT = Font(size=10, color="2D3748")
TOTAL_FONT = Font(size=11, bold=True, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin", color="D4C4A8"),
    right=Side(style="thin", color="D4C4A8"),
    top=Side(style="thin", color="D4C4A8"),
    bottom=Side(style="thin", color="D4C4A8"),
)
ROW_ALT_FILL = PatternFill(start_color="FBF8F3", end_color="FBF8F3", fill_type="solid")


def format_excel_datetime(value=None):
    value = value or datetime.now()
    return value.strftime("%d-%m-%Y %H:%M")


def format_excel_date(value):
    if not value:
        return ""
    return value.strftime("%d-%m-%Y")


def get_excel_branding_styles():
    return {
        "header_fill": HEADER_FILL,
        "subheader_fill": SUBHEADER_FILL,
        "white_font": WHITE_FONT,
        "title_font": TITLE_FONT,
        "subtitle_font": SUBTITLE_FONT,
        "meta_font": META_FONT,
        "data_font": DATA_FONT,
        "total_font": TOTAL_FONT,
        "border": THIN_BORDER,
        "row_alt_fill": ROW_ALT_FILL,
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center"),
        "right": Alignment(horizontal="right", vertical="center"),
        "wrap": Alignment(horizontal="left", vertical="center", wrap_text=True),
    }


def _add_company_logo(ws, empresa):
    logo = getattr(empresa, "logo", None)
    if not logo:
        return False

    try:
        logo_path = Path(logo.path)
    except Exception:
        return False

    if not logo_path.exists():
        return False

    try:
        image = OpenPyxlImage(str(logo_path))
        max_width = 180
        max_height = 70
        scale = min(max_width / image.width, max_height / image.height, 1)
        image.width = int(image.width * scale)
        image.height = int(image.height * scale)
        ws.add_image(image, "A1")
        return True
    except Exception:
        return False


def setup_corporate_sheet(ws, empresa, title, subtitle="", total_columns=1):
    styles = get_excel_branding_styles()
    total_columns = max(int(total_columns or 1), 1)
    last_col = get_column_letter(total_columns)
    has_logo = _add_company_logo(ws, empresa)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"

    title_start = "B" if has_logo and total_columns > 1 else "A"
    if has_logo:
        ws.column_dimensions["A"].width = 20

    ws.merge_cells(f"{title_start}1:{last_col}2")
    title_cell = ws[f"{title_start}1"]
    title_cell.value = title
    title_cell.font = styles["title_font"]
    title_cell.alignment = styles["center"]
    title_cell.fill = styles["subheader_fill"]

    ws.merge_cells(f"A3:{last_col}3")
    subtitle_cell = ws["A3"]
    subtitle_cell.value = subtitle or (getattr(empresa, "nombre", "") or "").upper()
    subtitle_cell.font = styles["subtitle_font"]
    subtitle_cell.alignment = styles["center"]

    ws.merge_cells(f"A4:{last_col}4")
    meta_cell = ws["A4"]
    rut = getattr(empresa, "rut", "") or "-"
    meta_cell.value = f"Empresa: {empresa.nombre} | RUT: {rut} | Exportado: {format_excel_datetime()}"
    meta_cell.font = styles["meta_font"]
    meta_cell.alignment = styles["center"]

    ws.merge_cells(f"A5:{last_col}5")
    band_cell = ws["A5"]
    band_cell.fill = styles["subheader_fill"]

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 8

    return {
        "header_row": 6,
        "data_row": 7,
        "styles": styles,
    }


def write_table_headers(ws, row, headers, styles=None):
    styles = styles or get_excel_branding_styles()
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.fill = styles["header_fill"]
        cell.font = styles["white_font"]
        cell.border = styles["border"]
        cell.alignment = styles["center"]
    ws.row_dimensions[row].height = 24


def style_data_row(ws, row, total_columns, styles=None, zebra=False):
    styles = styles or get_excel_branding_styles()
    fill = styles["row_alt_fill"] if zebra else None
    for col_idx in range(1, total_columns + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.border = styles["border"]
        if not cell.font or cell.font == Font():
            cell.font = styles["data_font"]
        if fill:
            cell.fill = fill


def set_column_widths(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def finalize_excel_table(ws, header_row, data_start_row, data_end_row, total_columns):
    if data_end_row >= data_start_row:
        last_col = get_column_letter(total_columns)
        ws.auto_filter.ref = f"A{header_row}:{last_col}{data_end_row}"


def build_excel_response(workbook, filename):
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response

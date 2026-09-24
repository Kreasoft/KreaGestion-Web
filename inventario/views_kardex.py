"""
Vistas para el Kardex (Cartola) de Artículos
"""
from django.shortcuts import render
from django.contrib.auth.decorators import login_required, permission_required
from django.db.models import Q
from django.utils import timezone
from django.http import HttpResponse
from datetime import datetime, timedelta
from core.decorators import requiere_empresa
from core.excel_utils import (
    build_excel_response,
    format_excel_date,
    get_excel_branding_styles,
    set_column_widths,
    setup_corporate_sheet,
    style_data_row,
    write_table_headers,
)
from articulos.models import Articulo
from bodegas.models import Bodega
from .models import Inventario, Stock
from openpyxl import Workbook
from openpyxl.styles import Font


def _obtener_datos_kardex(request):
    """Construye el contexto de datos del Kardex para vista y exportacion."""
    articulo_id = request.GET.get('articulo', '').strip()
    bodega_id = request.GET.get('bodega', '').strip()
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')

    if articulo_id:
        articulo_id = articulo_id.replace('\xa0', '').replace(' ', '').replace(',', '').replace('.', '')
    if bodega_id:
        bodega_id = bodega_id.replace('\xa0', '').replace(' ', '').replace(',', '').replace('.', '')

    if not fecha_desde:
        fecha_desde = (timezone.now() - timedelta(days=30)).strftime('%Y-%m-%d')
    if not fecha_hasta:
        fecha_hasta = timezone.now().strftime('%Y-%m-%d')

    articulo = None
    if articulo_id:
        try:
            articulo = Articulo.objects.get(id=int(articulo_id), empresa=request.empresa)
        except (Articulo.DoesNotExist, ValueError):
            articulo = None

    bodega = None
    bodegas = Bodega.objects.filter(empresa=request.empresa, activa=True).select_related('sucursal')
    if bodega_id:
        try:
            bodega = bodegas.get(id=int(bodega_id))
        except (Bodega.DoesNotExist, ValueError):
            bodega = None

    movimientos = []
    saldo_actual = 0

    if articulo:
        try:
            fecha_desde_obj = datetime.strptime(fecha_desde, '%Y-%m-%d').date()
            fecha_hasta_obj = datetime.strptime(fecha_hasta, '%Y-%m-%d').date()
        except Exception:
            fecha_desde_obj = (timezone.now() - timedelta(days=30)).date()
            fecha_hasta_obj = timezone.now().date()

        saldo_inicial = 0
        inventarios_iniciales = Inventario.objects.filter(
            empresa=request.empresa,
            articulo=articulo,
            fecha_movimiento__date__lt=fecha_desde_obj,
            estado='confirmado'
        )

        if bodega:
            inventarios_iniciales = inventarios_iniciales.filter(
                Q(bodega_origen=bodega) | Q(bodega_destino=bodega)
            )

        for inv in inventarios_iniciales:
            tipo_mov = inv.tipo_movimiento.lower()
            if tipo_mov == 'entrada':
                if not bodega or inv.bodega_destino == bodega:
                    saldo_inicial += inv.cantidad
            elif tipo_mov == 'salida':
                if not bodega or inv.bodega_origen == bodega:
                    saldo_inicial -= inv.cantidad
            elif tipo_mov == 'ajuste':
                if not bodega or inv.bodega_destino == bodega:
                    saldo_inicial += inv.cantidad
            elif tipo_mov == 'transferencia':
                if bodega:
                    if inv.bodega_origen == bodega:
                        saldo_inicial -= inv.cantidad
                    elif inv.bodega_destino == bodega:
                        saldo_inicial += inv.cantidad

        inventarios = Inventario.objects.filter(
            empresa=request.empresa,
            articulo=articulo,
            fecha_movimiento__date__gte=fecha_desde_obj,
            fecha_movimiento__date__lte=fecha_hasta_obj,
            estado='confirmado'
        ).select_related('bodega_origen', 'bodega_destino', 'creado_por')

        if bodega:
            inventarios = inventarios.filter(
                Q(bodega_origen=bodega) | Q(bodega_destino=bodega)
            )

        inventarios = inventarios.order_by('fecha_movimiento', 'id')

        movimientos_unicos = {}
        inventarios_limpios = []
        for inv in inventarios:
            clave = (
                inv.numero_documento or inv.numero_folio or '',
                inv.tipo_movimiento,
                inv.bodega_origen_id if inv.bodega_origen else None,
                inv.bodega_destino_id if inv.bodega_destino else None,
                str(inv.fecha_movimiento.date()),
                str(inv.cantidad),
            )
            if clave not in movimientos_unicos:
                movimientos_unicos[clave] = inv
                inventarios_limpios.append(inv)
            else:
                existente = movimientos_unicos[clave]
                if inv.id < existente.id:
                    index = inventarios_limpios.index(existente)
                    inventarios_limpios[index] = inv
                    movimientos_unicos[clave] = inv

        saldo = saldo_inicial
        for inv in inventarios_limpios:
            entrada = 0
            salida = 0
            tipo_mov = inv.tipo_movimiento.lower()

            if tipo_mov == 'entrada':
                if not bodega or inv.bodega_destino == bodega:
                    entrada = inv.cantidad
                    saldo += entrada
            elif tipo_mov == 'salida':
                if not bodega or inv.bodega_origen == bodega:
                    salida = inv.cantidad
                    saldo -= salida
            elif tipo_mov == 'ajuste':
                if not bodega or inv.bodega_destino == bodega:
                    if inv.cantidad >= 0:
                        entrada = inv.cantidad
                        saldo += entrada
                    else:
                        salida = abs(inv.cantidad)
                        saldo -= salida
            elif tipo_mov == 'transferencia':
                if bodega:
                    if inv.bodega_origen == bodega:
                        salida = inv.cantidad
                        saldo -= salida
                    elif inv.bodega_destino == bodega:
                        entrada = inv.cantidad
                        saldo += entrada
                else:
                    entrada = inv.cantidad
                    salida = inv.cantidad

            if entrada > 0 or salida > 0:
                movimientos.append({
                    'fecha': inv.fecha_movimiento,
                    'tipo': inv.get_tipo_movimiento_display(),
                    'tipo_codigo': inv.tipo_movimiento,
                    'documento': inv.numero_documento or inv.numero_folio or '-',
                    'detalle': inv.descripcion or inv.motivo or '-',
                    'bodega_origen': inv.bodega_origen.nombre if inv.bodega_origen else '-',
                    'bodega_destino': inv.bodega_destino.nombre if inv.bodega_destino else '-',
                    'entrada': entrada,
                    'salida': salida,
                    'saldo': saldo,
                    'usuario': inv.creado_por.username if inv.creado_por else '-'
                })
        saldo_actual = saldo

    articulos = Articulo.objects.filter(
        empresa=request.empresa,
        activo=True
    ).order_by('nombre')

    stock_sistema = None
    if articulo:
        stock_sistema_calc = 0
        inv_qs = Inventario.objects.filter(
            empresa=request.empresa,
            articulo=articulo,
            estado='confirmado'
        ).select_related('bodega_origen', 'bodega_destino')
        try:
            fecha_hasta_obj
        except NameError:
            fecha_hasta_obj = timezone.now().date()
        inv_qs = inv_qs.filter(fecha_movimiento__date__lte=fecha_hasta_obj)
        if bodega:
            inv_qs = inv_qs.filter(Q(bodega_origen=bodega) | Q(bodega_destino=bodega))

        movimientos_unicos_stock = {}
        inv_qs_limpios = []
        for inv in inv_qs:
            clave = (
                inv.numero_documento or inv.numero_folio or '',
                inv.tipo_movimiento,
                inv.bodega_origen_id if inv.bodega_origen else None,
                inv.bodega_destino_id if inv.bodega_destino else None,
                str(inv.fecha_movimiento.date()),
                str(inv.cantidad),
            )
            if clave not in movimientos_unicos_stock:
                movimientos_unicos_stock[clave] = inv
                inv_qs_limpios.append(inv)
            else:
                existente = movimientos_unicos_stock[clave]
                if inv.id < existente.id:
                    index = inv_qs_limpios.index(existente)
                    inv_qs_limpios[index] = inv
                    movimientos_unicos_stock[clave] = inv

        for inv in inv_qs_limpios:
            t = (inv.tipo_movimiento or '').lower()
            if t == 'entrada':
                if not bodega or inv.bodega_destino == bodega:
                    stock_sistema_calc += inv.cantidad
            elif t == 'salida':
                if not bodega or inv.bodega_origen == bodega:
                    stock_sistema_calc -= inv.cantidad
            elif t == 'ajuste':
                if not bodega or inv.bodega_destino == bodega:
                    stock_sistema_calc += inv.cantidad
            elif t == 'transferencia':
                if bodega:
                    if inv.bodega_origen == bodega:
                        stock_sistema_calc -= inv.cantidad
                    elif inv.bodega_destino == bodega:
                        stock_sistema_calc += inv.cantidad
        stock_sistema = stock_sistema_calc

    return {
        'articulo': articulo,
        'bodega': bodega,
        'bodegas': bodegas,
        'articulos': articulos,
        'movimientos': movimientos,
        'saldo_actual': saldo_actual,
        'stock_sistema': stock_sistema,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'tiene_sucursal': hasattr(request, 'sucursal_activa') and request.sucursal_activa,
        'sucursal_nombre': request.sucursal_activa.nombre if hasattr(request, 'sucursal_activa') and request.sucursal_activa else None,
    }


@login_required
@requiere_empresa
@permission_required('inventario.view_inventario', raise_exception=True)
def kardex_articulo(request):
    """
    Kardex (Cartola) de movimientos de un artículo
    Muestra todas las entradas, salidas, transferencias y ajustes
    """
    context = _obtener_datos_kardex(request)
    return render(request, 'inventario/kardex_articulo.html', context)


@login_required
@requiere_empresa
@permission_required('inventario.view_inventario', raise_exception=True)
def kardex_articulo_excel(request):
    """Exporta el Kardex de articulo a un Excel real con formato corporativo."""
    context = _obtener_datos_kardex(request)
    articulo = context['articulo']
    if not articulo:
        return HttpResponse("Debe seleccionar un articulo valido para exportar el Kardex.", status=400)

    movimientos = context['movimientos']
    bodega = context['bodega']
    fecha_desde = context['fecha_desde']
    fecha_hasta = context['fecha_hasta']
    saldo_actual = context['saldo_actual']

    wb = Workbook()
    ws = wb.active
    ws.title = "Kardex"

    headers = ['Fecha', 'Tipo', 'Documento', 'Detalle', 'Origen', 'Destino', 'Entrada', 'Salida', 'Saldo', 'Usuario']
    branding = setup_corporate_sheet(
        ws,
        request.empresa,
        "KARDEX DE ARTICULO",
        f"Articulo: {articulo.codigo} - {articulo.nombre}",
        len(headers),
    )
    styles = branding['styles']
    ws['A5'] = (
        f"Bodega: {(bodega.nombre if bodega else 'Todas')} | "
        f"Periodo: {fecha_desde} al {fecha_hasta} | "
        f"Saldo actual: {saldo_actual}"
    )
    ws['A5'].font = styles['subtitle_font']
    ws['A5'].alignment = styles['center']

    write_table_headers(ws, branding['header_row'], headers, styles)

    row = branding['data_row']
    for movimiento in movimientos:
        ws.cell(row=row, column=1, value=movimiento['fecha'].strftime('%d-%m-%Y %H:%M')).alignment = styles['center']
        ws.cell(row=row, column=2, value=movimiento['tipo']).alignment = styles['center']
        ws.cell(row=row, column=3, value=movimiento['documento']).alignment = styles['center']
        ws.cell(row=row, column=4, value=movimiento['detalle']).alignment = styles['wrap']
        ws.cell(row=row, column=5, value=movimiento['bodega_origen'])
        ws.cell(row=row, column=6, value=movimiento['bodega_destino'])
        ws.cell(row=row, column=7, value=float(movimiento['entrada'] or 0)).alignment = styles['right']
        ws.cell(row=row, column=8, value=float(movimiento['salida'] or 0)).alignment = styles['right']
        ws.cell(row=row, column=9, value=float(movimiento['saldo'] or 0)).alignment = styles['right']
        ws.cell(row=row, column=10, value=movimiento['usuario']).alignment = styles['center']

        for col in (7, 8, 9):
            ws.cell(row=row, column=col).number_format = '#,##0.00'
        style_data_row(ws, row, len(headers), styles, zebra=((row - branding['data_row']) % 2 == 1))
        row += 1

    resumen_row = row
    ws.cell(row=resumen_row, column=1, value='RESUMEN').font = Font(bold=True)
    ws.cell(row=resumen_row, column=8, value='Saldo final').font = Font(bold=True)
    ws.cell(row=resumen_row, column=8).alignment = styles['right']
    ws.cell(row=resumen_row, column=9, value=float(saldo_actual or 0)).font = Font(bold=True)
    ws.cell(row=resumen_row, column=9).alignment = styles['right']
    ws.cell(row=resumen_row, column=9).number_format = '#,##0.00'
    style_data_row(ws, resumen_row, len(headers), styles)

    set_column_widths(ws, [18, 14, 16, 44, 22, 22, 12, 12, 12, 16])
    ws.freeze_panes = f"A{branding['data_row']}"
    safe_codigo = ''.join(ch for ch in articulo.codigo if ch.isalnum() or ch in ('-', '_')) or 'articulo'
    return build_excel_response(wb, f"kardex_{safe_codigo}.xlsx")
